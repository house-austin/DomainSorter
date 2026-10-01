import os

from openpyxl import Workbook, load_workbook


class Domain:
    def __init__(self, domain):
        self.domain = domain

    def get_sld(self):
        """Second-level domain, e.g. 'mail.example.com' -> 'example'."""
        parts = self.domain.split(".")
        # Fall back to the whole string if there's no TLD to strip
        return parts[-2] if len(parts) >= 2 else self.domain


def sorted_path(path):
    """'domains.txt' -> 'domains_sorted.txt'"""
    base, ext = os.path.splitext(path)
    return f"{base}_sorted{ext}"


def read_txt(path):
    with open(path, "r", encoding="utf-8") as f:
        return [line.strip() for line in f if line.strip()]


def read_xlsx(path):
    """Prompt for a sheet name, then read column A of that sheet."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        print("Sheets found: " + ", ".join(wb.sheetnames))
        while True:
            sheet_name = input("Please enter the sheet name: ").strip()
            if sheet_name in wb.sheetnames:
                break
            print("Sheet not found. Please try again.")

        ws = wb[sheet_name]
        domains = []
        for (cell,) in ws.iter_rows(min_col=1, max_col=1, values_only=True):
            if cell is not None and str(cell).strip():
                domains.append(str(cell).strip())
        return domains, sheet_name
    finally:
        wb.close()


def read_file():
    """Keep prompting until a readable, non-empty file is loaded."""
    while True:
        path = input("Please enter filename (Make sure to include file extension): ").strip()
        try:
            if path.lower().endswith(".xlsx"):
                domains, sheet_name = read_xlsx(path)
            else:
                domains, sheet_name = read_txt(path), None

            if not domains:
                print("No domains found in file. Please try again.")
                continue

            print("File loaded!")
            return [Domain(d) for d in domains], path, sheet_name

        except (FileNotFoundError, OSError):
            print("Could not read file. Please try again")
        except Exception as e:  # e.g. corrupt/invalid .xlsx
            print(f"Could not read file ({e}). Please try again")


def write_file(domain_list, input_path, sheet_name=None):
    out_path = sorted_path(input_path)
    try:
        if input_path.lower().endswith(".xlsx"):
            wb = Workbook()
            ws = wb.active
            ws.title = sheet_name or "Sheet1"
            for row, domain in enumerate(domain_list, start=1):
                ws.cell(row=row, column=1, value=domain.domain)
            wb.save(out_path)
        else:
            with open(out_path, "w", encoding="utf-8") as f:
                for domain in domain_list:
                    f.write(domain.domain + "\n")
        print(f"Created file {out_path}")
    except OSError:
        print("Could not write file.")


def sort_domains(domain_list):
    # Original: bubble sort alphabetically, then a stable bubble sort by SLD.
    # Python's sort is stable, so one sort on (SLD, full domain) is equivalent.
    domain_list.sort(key=lambda d: (d.get_sld().lower(), d.domain.lower()))


def main():
    print("Welcome! This program sorts domains based on the Second Level Domain! "
          "Please begin by supplying a file containing domains and subdomains.")
    domain_list, input_path, sheet_name = read_file()
    sort_domains(domain_list)
    print("Domains sorted!")
    write_file(domain_list, input_path, sheet_name)


if __name__ == "__main__":
    main()